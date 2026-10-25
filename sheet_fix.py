#!/usr/bin/env python3
"""
SheetFix engine (v1) — audit -> fix -> re-check for Excel workbooks.

Wraps spreadsheet-auditor (MIT, petehottelet) for detection and re-check.
The fix layer is ours: it applies only mechanical, evidence-backed fixes and
reports everything judgment-based with the auditor's own suggestion.

Usage: python3 sheet_fix.py <workbook.xlsx> <outdir> [auditor_cli]
Outputs in outdir: <name>_fixed.xlsx, fix_report.md, findings_before.json, findings_after.json
"""
import json, os, re, subprocess, sys

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter

FIXABLE_ORDER = ["NUMBERS_STORED_AS_TEXT", "WHITESPACE_KEY", "FORMULA_DRIFT", "RANGE_EXCLUSION"]
REF_RE = re.compile(r"(\$?)([A-Z]{1,3})(\$?)(\d+)")


def run_auditor(cli, path, out_json):
    try:
        from spreadsheet_auditor.audit import main as auditor_main
        auditor_main([path, "--json", out_json, "--out", out_json + ".md"])
    except ImportError:
        subprocess.run([cli, path, "--json", out_json, "--out", out_json + ".md"],
                       capture_output=True, text=True)  # exit code >0 just means findings exist
    with open(out_json, encoding="utf-8") as fh:
        return json.load(fh)["findings"]


def translate_formula(formula, from_cell, to_cell):
    """Shift relative refs in `formula` as if copied from from_cell to to_cell (A1 style)."""
    fc = re.fullmatch(r"([A-Z]{1,3})(\d+)", from_cell)
    tc = re.fullmatch(r"([A-Z]{1,3})(\d+)", to_cell)
    dcol = column_index_from_string(tc.group(1)) - column_index_from_string(fc.group(1))
    drow = int(tc.group(2)) - int(fc.group(2))

    def repl(m):
        cold, col, rowd, row = m.groups()
        ncol = col if cold else get_column_letter(column_index_from_string(col) + dcol)
        nrow = row if rowd else str(int(row) + drow)
        return f"{cold}{ncol}{rowd}{nrow}"

    return REF_RE.sub(repl, formula)


def neighbor_formula(finding):
    """Extract (cell, formula) from a FORMULA_DRIFT evidence line like 'Budget!C10==SUM(C7:C9)'."""
    for ev in finding.get("evidence") or []:
        m = re.search(r"([A-Z]{1,3}\d+)=(=[^;]+)", ev.strip())
        if m:
            return m.group(1), m.group(2).strip()
    return None, None


def apply_fixes(wb, findings):
    fixes, skipped = [], []
    by_cell_fixed = set()
    ordered = sorted(findings, key=lambda f: FIXABLE_ORDER.index(f["rule_id"]) if f["rule_id"] in FIXABLE_ORDER else 99)
    for f in ordered:
        rule, sheet, cell = f["rule_id"], f["sheet"], f["cell"]
        if rule not in FIXABLE_ORDER or sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        target = ws[cell]
        old = target.value

        if rule == "NUMBERS_STORED_AS_TEXT":
            if isinstance(old, str):
                txt = old.replace(",", "").replace(" ", "").strip()
                try:
                    new = float(txt)
                    new = int(new) if new == int(new) else new
                except ValueError:
                    skipped.append((f, "value does not parse as a number"))
                    continue
                target.value = new
                fixes.append((f, old, new, "text number converted to a real number"))
                by_cell_fixed.add((sheet, cell))

        elif rule == "WHITESPACE_KEY":
            if isinstance(old, str) and old != old.strip():
                target.value = old.strip()
                fixes.append((f, old, target.value, "leading/trailing spaces trimmed"))
                by_cell_fixed.add((sheet, cell))

        elif rule == "FORMULA_DRIFT":
            ncell, nformula = neighbor_formula(f)
            if not ncell or not isinstance(old, str) or not old.startswith("="):
                skipped.append((f, "no neighbor pattern available to translate"))
                continue
            new = translate_formula(nformula, ncell, cell)
            if new == old:
                skipped.append((f, "translated pattern equals current formula"))
                continue
            target.value = new
            fixes.append((f, old, new, f"formula restored to the block's dominant pattern (translated from {ncell})"))
            by_cell_fixed.add((sheet, cell))

        elif rule == "RANGE_EXCLUSION":
            if (sheet, cell) in by_cell_fixed:
                continue  # already repaired by the drift fix on the same cell
            ev = " ".join(f.get("evidence") or [])
            m = re.search(r"([A-Z]{1,3}\d+:\d+|[A-Z]{1,3}\d+:[A-Z]{1,3}\d+)", str(old or ""))
            ex = re.search(r"stops short of (?:[A-Za-z_0-9]+!)?([A-Z]{1,3}\d+)", ev)
            if not (isinstance(old, str) and old.startswith("=") and ex):
                skipped.append((f, "could not identify the excluded cell from the evidence"))
                continue
            excl = ex.group(1)
            rm = re.fullmatch(r"=SUM\(([A-Z]{1,3})(\d+):([A-Z]{1,3})(\d+)\)", old.replace(" ", ""), re.I)
            if not rm:
                skipped.append((f, "range is not a simple SUM range; left for manual review"))
                continue
            c1, r1, c2, r2 = rm.group(1), int(rm.group(2)), rm.group(3), int(rm.group(4))
            ec, er = re.fullmatch(r"([A-Z]{1,3})(\d+)", excl).groups()
            er = int(er)
            if ec == c1 == c2 and er == r2 + 1:
                new = f"=SUM({c1}{r1}:{c2}{er})"
            elif ec == c1 == c2 and er == r1 - 1:
                new = f"=SUM({c1}{er}:{c2}{r2})"
            elif er == r1 == r2 and column_index_from_string(ec) == column_index_from_string(c2) + 1:
                new = f"=SUM({c1}{r1}:{ec}{r2})"
            else:
                skipped.append((f, "excluded cell is not directly adjacent to the range"))
                continue
            target.value = new
            fixes.append((f, old, new, f"range extended to include the excluded cell {excl}"))
            by_cell_fixed.add((sheet, cell))
    return fixes, skipped


def key(f):
    return (f["rule_id"], f["location"])


def run(src, outdir, cli="spreadsheet-auditor"):
    """Full audit -> fix -> re-check pipeline. Returns a summary dict (used by CLI and GUI)."""
    os.makedirs(outdir, exist_ok=True)
    stem = os.path.splitext(os.path.basename(src))[0]

    before = run_auditor(cli, src, os.path.join(outdir, "findings_before.json"))
    wb = openpyxl.load_workbook(src)
    fixes, skipped = apply_fixes(wb, before)
    fixed_path = os.path.join(outdir, f"{stem}_fixed.xlsx")
    wb.save(fixed_path)
    after = run_auditor(cli, fixed_path, os.path.join(outdir, "findings_after.json"))

    before_keys = {key(f) for f in before}
    after_keys = {key(f) for f in after}
    resolved = [f for f in before if key(f) not in after_keys]
    remaining = [f for f in after if key(f) in before_keys]
    new_issues = [f for f in after if key(f) not in before_keys]

    lines = ["# SheetFix report — audit, fix & re-check", ""]
    lines.append(f"Workbook: {os.path.basename(src)}")
    lines.append(f"- Findings before: {len(before)}  ·  auto-fixed: {len(fixes)}  ·  resolved on re-check: {len(resolved)}  ·  remaining: {len(remaining)}  ·  new issues introduced: {len(new_issues)}")
    lines.append("")
    lines.append("## Fixes applied (each one re-verified by re-running the audit)")
    if fixes:
        for f, old, new, basis in fixes:
            lines.append(f"- [{f['location']}] {f['rule_id']}: `{old}` -> `{new}` ({basis})")
    else:
        lines.append("- none")
    if skipped:
        lines.append("")
        lines.append("## Fixable-looking findings left untouched")
        for f, why in skipped:
            lines.append(f"- [{f['location']}] {f['rule_id']}: {why}")
    lines.append("")
    lines.append("## Remaining findings — need your judgment, not a mechanical fix")
    if remaining:
        for f in remaining:
            lines.append(f"- [{f['severity']}] [{f['location']}] {f['rule_id']}: {f['title']}. Suggestion: {f.get('suggested_fix') or 'review manually'}")
    else:
        lines.append("- none — the fixed workbook passes the full audit set")
    if new_issues:
        lines.append("")
        lines.append("## WARNING: findings that appeared only after fixing")
        for f in new_issues:
            lines.append(f"- [{f['severity']}] [{f['location']}] {f['rule_id']}: {f['title']}")
    lines.append("")
    lines.append("Scope note: fixes are mechanical and evidence-backed (text numbers, padded keys, drifted formulas restored to the block pattern, sum ranges missing an adjacent row). Anything that changes business meaning is reported, never silently altered. All processing is local; the workbook never leaves this machine.")
    report_path = os.path.join(outdir, "fix_report.md")
    with open(report_path, "w", encoding="utf-8") as fh:
        fh.write("\n".join(lines) + "\n")
    return {"before": before, "fixes": fixes, "skipped": skipped, "resolved": resolved,
            "remaining": remaining, "new_issues": new_issues,
            "fixed_path": fixed_path, "report_path": report_path}


def main():
    src, outdir = sys.argv[1], sys.argv[2]
    cli = sys.argv[3] if len(sys.argv) > 3 else "spreadsheet-auditor"
    s = run(src, outdir, cli)
    print(f"before={len(s['before'])} fixes={len(s['fixes'])} resolved={len(s['resolved'])} remaining={len(s['remaining'])} new={len(s['new_issues'])}")


if __name__ == "__main__":
    main()
